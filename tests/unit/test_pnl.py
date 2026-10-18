"""Monthly P&L: settlement file readers and PnlService.

Synthetic files reproduce the real 2026 report layouts (Amazon preamble with
a decoy header line and comma-formatted amounts; Flipkart / Meesho two-row
headers with a formula row), so no business data is stored in the repo.
"""
import io

import openpyxl
import pandas as pd
import pytest
from sqlalchemy import create_engine, text

from src.infrastructure import init_db
from src.infrastructure.logger import DataValidationException
from src.infrastructure.settlement_parsers import (
    detect_platform, parse_amazon, parse_flipkart, parse_meesho,
)

# ---------------------------------------------------------------------------
# Synthetic reports
# ---------------------------------------------------------------------------

AMAZON_HEADER = [
    "date/time", "settlement id", "type", "order id", "Sku", "description", "quantity",
    "marketplace", "account type", "fulfillment", "order city", "order state", "order postal",
    "product sales", "shipping credits", "gift wrap credits", "promotional rebates",
    "Total sales tax liable(GST before adjusting TCS)", "TCS-CGST", "TCS-SGST", "TCS-IGST",
    "TDS (Section 194-O)", "selling fees", "fba fees", "other transaction fees", "other",
    "total", "Transaction Status", "Transaction Release Date",
]


def _amazon_row(when, ttype, order, sku, desc, qty, sales=0, tax=0, tcs=0, tds=0,
                selling=0, fba=0, other_txn=0, other=0, status="Released"):
    total = sales + tax + tcs + tds + selling + fba + other_txn + other
    vals = [when, "1", ttype, order, sku, desc, qty, "amazon.in", "x", "Amazon", "", "", "",
            sales, 0, 0, 0, tax, tcs, 0, 0, tds, selling, fba, other_txn, other,
            total, status, ""]
    out = []
    for v in vals:
        if isinstance(v, (int, float)):
            out.append(f'"{v:,.2f}"')  # Amazon quotes amounts and uses thousands commas
        else:
            out.append(f'"{v}"')
    return ",".join(out)


def make_amazon_csv():
    preamble = [
        '"Includes Amazon Marketplace, Fulfillment by Amazon (FBA), and Amazon Webstore transactions"',
        '"All amounts in INR, unless specified"',
        '"Date/Time range: 1 Sep 2026 - 30 Sep 2026"',  # decoy: starts with date/time
    ]
    header = ",".join(f'"{h}"' for h in AMAZON_HEADER)
    rows = [
        # 31 Aug 19:00 UTC is 1 Sep 00:30 IST -> September
        _amazon_row("31 Aug 2026 7:00:00 pm UTC", "Order", "A-1", "CONT-777-PO2", "box", 2,
                    sales=1200.0, tax=216.0, tcs=-6.0, tds=-1.2, selling=-118.0, fba=-236.0),
        _amazon_row("5 Sept 2026 10:00:00 am UTC", "Refund", "A-2", "CONT-777-PO2", "box", 1,
                    sales=-1016.10, tax=-182.90, tcs=5.08, selling=23.46),
        _amazon_row("6 Sept 2026 10:00:00 am UTC", "Transfer", "", "", "payout", "",
                    other=-5000.0),
        _amazon_row("7 Sept 2026 10:00:00 am UTC", "Shipping Services", "A-1", "", "Easy Ship", "",
                    other_txn=-64.90),
        _amazon_row("8 Sept 2026 10:00:00 am UTC", "Service Fee", "", "", "Cost of Advertising", "",
                    other_txn=-1180.0),
        _amazon_row("9 Sept 2026 10:00:00 am UTC", "Fulfilment Fee Refund", "A-1", "",
                    "Fulfillment Fee Refund", "", fba=76.70),
        _amazon_row("10 Sept 2026 10:00:00 am UTC", "Reimbursements", "", "", "lost package", "",
                    other=500.0, status="Deferred"),
    ]
    return ("\n".join(preamble + [header] + rows) + "\n").encode("utf-8-sig")


def _two_row_header_sheet(wb, name, banner, columns, rows):
    ws = wb.create_sheet(name)
    ws.append([banner])
    ws.append(columns)
    ws.append(["=formula" for _ in columns])
    for r in rows:
        ws.append([r.get(c) for c in columns])


FK_ORDER_COLS = [
    "NEFT ID", "Payment Date", "Bank Settlement Value (Rs.) \n= SUM(J:R)", "Order ID",
    "Sale Amount (Rs.)", "Total Offer Amount (Rs.)", "My share (Rs.)",
    "Customer Add-ons Amount (Rs.)", "Marketplace Fee (Rs.) \n= SUM (V:AI)", "Taxes (Rs.)",
    "Offer Adjustments (Rs.)", "Protection Fund (Rs.)", "Refund (Rs.)",
    "Commission (Rs.)", "Fixed Fee  (Rs.)", "Shipping Fee (Rs.)", "Reverse Shipping Fee (Rs.)",
    "TCS (Rs.)", "TDS (Rs.)", "GST on MP Fees (Rs.)", "Item GST Rate (%)", "Seller SKU", "Quantity",
]


def _fk_order(order, sku, qty, sale=0.0, refund=0.0, fixed=0.0, reverse=0.0, tcs=0.0, tds=0.0):
    gst = round((fixed + reverse) * 0.18, 2)
    mp = fixed + reverse
    taxes = tcs + tds + gst
    bank = sale + refund + mp + taxes
    return {
        "NEFT ID": "N1", "Payment Date": "2026-09-10",
        "Bank Settlement Value (Rs.) \n= SUM(J:R)": bank, "Order ID": order,
        "Sale Amount (Rs.)": sale, "Total Offer Amount (Rs.)": 0, "My share (Rs.)": 0,
        "Customer Add-ons Amount (Rs.)": 0, "Marketplace Fee (Rs.) \n= SUM (V:AI)": mp,
        "Taxes (Rs.)": taxes, "Offer Adjustments (Rs.)": 0, "Protection Fund (Rs.)": 0,
        "Refund (Rs.)": refund, "Commission (Rs.)": 0, "Fixed Fee  (Rs.)": fixed,
        "Shipping Fee (Rs.)": 0, "Reverse Shipping Fee (Rs.)": reverse, "TCS (Rs.)": tcs,
        "TDS (Rs.)": tds, "GST on MP Fees (Rs.)": gst, "Item GST Rate (%)": 0,
        "Seller SKU": sku, "Quantity": qty,
    }


def make_flipkart_xlsx():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _two_row_header_sheet(wb, "Orders", "Payment Details", FK_ORDER_COLS, [
        _fk_order("OD1", "caps-para-white", 2, sale=420.0, fixed=-10.0, tcs=-2.0, tds=-0.4),
        _fk_order("OD2", "CAPS-PARA-WHITE", 1, refund=-210.0, reverse=-175.0),
        _fk_order("OD3", "CAPS-PARA-WHITE", 1),  # neither settled nor refunded: void
    ])
    _two_row_header_sheet(wb, "Ads", "Payment Details", [
        "NEFT ID", "Payment Date", "Settlement Value (Rs.) = SUM(G:K)", "Type",
        "Campaign / Transaction ID", "Wallet Redeem (Rs.)", "Wallet Topup (Rs.)",
    ], [{"NEFT ID": "N1", "Payment Date": "2026-09-12",
         "Settlement Value (Rs.) = SUM(G:K)": -118.0, "Type": "Topup",
         "Campaign / Transaction ID": "C1", "Wallet Redeem (Rs.)": 0, "Wallet Topup (Rs.)": -118.0}])
    _two_row_header_sheet(wb, "TCS_Recovery", "Payment Details", ["NEFT ID"], [])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MS_COLS = [
    "Sub Order No", "Supplier SKU", "Live Order Status", "Product GST %", "Quantity",
    "Payment Date", "Final Settlement Amount", "Total Sale Amount (Incl. Shipping & GST)",
    "Total Sale Return Amount (Incl. Shipping & GST)", "Meesho Commission (Incl. GST)",
    "Return Shipping Charge (Incl. GST)", "Shipping Charge (Incl. GST)", "TCS", "TDS",
]


def make_meesho_xlsx(ads_rows=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def row(sub, status, sale, ret, ship, rship):
        tcs, tds = -1.0, -0.2
        final = sale + ret + ship + rship + tcs + tds
        return {"Sub Order No": sub, "Supplier SKU": "L-CAP-RAIN-NAVY", "Live Order Status": status,
                "Product GST %": 5, "Quantity": 1, "Payment Date": "2026-09-04",
                "Final Settlement Amount": final,
                "Total Sale Amount (Incl. Shipping & GST)": sale,
                "Total Sale Return Amount (Incl. Shipping & GST)": ret,
                "Meesho Commission (Incl. GST)": 0, "Return Shipping Charge (Incl. GST)": rship,
                "Shipping Charge (Incl. GST)": ship, "TCS": tcs, "TDS": tds}

    _two_row_header_sheet(wb, "Order Payments", "Order Related Details", MS_COLS, [
        row("S1_1", "Delivered", 210.0, 0.0, -59.0, 0.0),
        row("S2_1", "RTO", 210.0, -210.0, -59.0, -29.5),
    ])
    ads = wb.create_sheet("Ads Cost")
    ads.append(["Ads Cost"])
    ads.append(["Deduction Duration", "Total Ads Cost"])
    ads.append(["=f", "=f"])
    for r in (ads_rows or [["No data is available for these dates."]]):
        ads.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _identity(rows):
    explained = (rows.sales_incl_gst + rows.fees_ex_gst + rows.shipping_ex_gst + rows.ads_ex_gst
                 + rows.other_ex_gst - rows.itc + rows.tcs + rows.tds)
    return (rows.settlement_value - explained).abs().max()


# ===========================================================================
# Readers
# ===========================================================================

@pytest.mark.unit
class TestDetect:

    def test_detects_each_platform_by_content(self):
        assert detect_platform(make_amazon_csv(), "x.csv") == "amazon"
        assert detect_platform(make_flipkart_xlsx(), "c463cdcf-a1a4.xlsx") == "flipkart"
        assert detect_platform(make_meesho_xlsx(), "anything.xlsx") == "meesho"

    def test_unknown_file(self):
        assert detect_platform(b"sku,qty\nA,1\n", "stock.csv") is None


@pytest.mark.unit
class TestAmazonReader:

    @pytest.fixture
    def rows(self):
        return parse_amazon(make_amazon_csv())

    def test_transfer_rows_are_excluded(self, rows):
        assert len(rows) == 6
        assert rows.settlement_value.sum() == pytest.approx(
            1200 + 216 - 6 - 1.2 - 118 - 236
            + (-1016.10 - 182.90 + 5.08 + 23.46) - 64.90 - 1180 + 76.70 + 500
        )

    def test_comma_amounts_are_read(self, rows):
        refund = rows[rows.line_type == "refund"].iloc[0]
        assert refund.sales_incl_gst == pytest.approx(-1199.0)
        assert refund.returned_units == 1

    def test_every_row_reconciles(self, rows):
        assert _identity(rows) < 0.01

    def test_utc_converted_to_india_time(self, rows):
        assert rows.iloc[0].txn_date == "2026-09-01"

    def test_fees_split_into_ex_gst_and_credit(self, rows):
        order = rows[rows.line_type == "order"].iloc[0]
        assert order.fees_ex_gst == pytest.approx(-100.0)
        assert order.shipping_ex_gst == pytest.approx(-200.0)
        assert order.itc == pytest.approx(18.0 + 36.0)
        assert order.gst_collected == pytest.approx(216.0)
        assert order.units == 2

    def test_non_order_rows_classified(self, rows):
        by_type = rows.groupby("line_type")
        assert by_type.ads_ex_gst.sum()["ads"] == pytest.approx(-1000.0)
        assert by_type.shipping_ex_gst.sum()["shipping"] == pytest.approx(-55.0 + 65.0)
        assert by_type.other_ex_gst.sum()["income"] == pytest.approx(500.0)
        assert rows[rows.line_type == "income"].is_deferred.iloc[0] == 1


@pytest.mark.unit
class TestFlipkartReader:

    @pytest.fixture
    def rows(self):
        return parse_flipkart(make_flipkart_xlsx())

    def test_void_lines_skipped_and_skus_uppercased(self, rows):
        orders = rows[rows.line_type == "order"]
        assert len(orders) == 2
        assert set(orders.sku_raw) == {"CAPS-PARA-WHITE"}

    def test_units_and_returns(self, rows):
        orders = rows[rows.line_type == "order"]
        assert orders.units.sum() == 2
        assert orders.returned_units.sum() == 1

    def test_actual_gst_on_fees_used(self, rows):
        orders = rows[rows.line_type == "order"]
        assert orders.itc.sum() == pytest.approx(1.8 + 31.5)
        assert orders.shipping_ex_gst.sum() == pytest.approx(-175.0)
        assert orders.gst_collected.isna().all()

    def test_ads_from_wallet_topup(self, rows):
        ads = rows[rows.line_type == "ads"].iloc[0]
        assert ads.ads_ex_gst == pytest.approx(-100.0)
        assert ads.itc == pytest.approx(18.0)

    def test_every_row_reconciles(self, rows):
        assert _identity(rows) < 0.01


@pytest.mark.unit
class TestMeeshoReader:

    def test_rto_counts_as_returned_and_rows_reconcile(self):
        rows = parse_meesho(make_meesho_xlsx())
        assert rows.units.sum() == 2
        assert rows.returned_units.sum() == 1
        assert rows.gst_collected.sum() == pytest.approx(210 * 5 / 105)
        assert rows.shipping_ex_gst.sum() == pytest.approx((-59 - 59 - 29.5) / 1.18)
        assert _identity(rows) < 0.01

    def test_unknown_ads_rows_are_refused_not_guessed(self):
        with pytest.raises(DataValidationException, match="Ads Cost"):
            parse_meesho(make_meesho_xlsx(ads_rows=[["Sep 2026", 500]]))


# ===========================================================================
# Service
# ===========================================================================

@pytest.fixture
def engine(tmp_path, monkeypatch):
    eng = create_engine(f"sqlite:///{tmp_path / 'pnl.db'}")
    with open(init_db.SCHEMA_PATH, encoding="utf-8") as f:
        with eng.connect() as conn:
            for stmt in (s for s in f.read().split(";") if s.strip()):
                conn.execute(text(stmt))
            conn.commit()
    init_db._apply_table_migrations(eng)
    init_db._apply_column_migrations(eng)
    init_db._apply_account_schema(eng)
    with eng.begin() as conn:
        conn.execute(text("PRAGMA foreign_keys = OFF"))
        conn.execute(text("INSERT INTO seller_accounts (account_id, display_name) VALUES "
                          "('account1','Account 1'), ('account2','Account 2')"))
        conn.execute(text("INSERT INTO config (marketplace) VALUES ('Amazon'), ('Flipkart'), ('Meesho')"))
        conn.execute(text("INSERT INTO product_master (sku, total_unit_cogs, gst_rate) VALUES "
                          "('CONT-777', 100, 18), ('CAP-PARA-WHITE', 50, 5)"))
        conn.execute(text("INSERT INTO pack_master (pack_sku, master_sku, quantity, packaging_cogs) VALUES "
                          "('CONT-777-PO2', 'CONT-777', 2, 10), ('CAP-PARA-WHITE-PK1', 'CAP-PARA-WHITE', 1, 3)"))
        conn.execute(text("INSERT INTO channel_listings (channel_sku, marketplace, internal_sku, account_id) VALUES "
                          "('CAPS-PARA-WHITE', 'Amazon', 'CAP-PARA-WHITE-PK1', 'account2')"))
    for mod in ("src.core.services.pnl_service", "src.core.services.catalog_service"):
        monkeypatch.setattr(f"{mod}.get_engine", lambda: eng)
    return eng


@pytest.fixture
def service(engine):
    from src.core.services.pnl_service import PnlService
    return PnlService()


@pytest.mark.unit
class TestImports:

    def test_import_then_same_file_refused(self, service):
        r = service.import_file(make_amazon_csv(), "a.csv", "account2")
        assert r["row_count"] == 6 and r["date_from"] == "2026-09-01"
        with pytest.raises(DataValidationException, match="already imported"):
            service.import_file(make_amazon_csv(), "copy.csv", "account2")

    def test_overlapping_period_refused(self, service):
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        redownloaded = make_amazon_csv().replace(b"lost package", b"lost parcel")
        with pytest.raises(DataValidationException, match="overlaps"):
            service.import_file(redownloaded, "a2.csv", "account2")

    def test_same_period_other_account_allowed(self, service):
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        other = make_amazon_csv().replace(b"lost package", b"lost parcel")
        service.import_file(other, "b.csv", "account1")
        assert len(service.list_imports()) == 2

    def test_unknown_account_refused(self, service):
        with pytest.raises(DataValidationException, match="does not exist"):
            service.import_file(make_amazon_csv(), "a.csv", "account9")

    def test_delete_removes_lines(self, service, engine):
        r = service.import_file(make_amazon_csv(), "a.csv", "account2")
        service.delete_import(r["import_id"])
        with engine.connect() as conn:
            assert conn.execute(text("SELECT COUNT(*) FROM settlement_lines")).scalar() == 0
        assert service.available_months() == []


@pytest.mark.unit
class TestMonthlyPnl:

    def test_cogs_on_kept_units_packaging_on_shipped(self, service):
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        res = service.monthly_pnl("2026-09")
        s = res["summary"]["Total"]
        # 2 packs sold, 1 refunded: kept 1 pack x 2 units x Rs100
        assert s["Goods cost (units kept)"] == pytest.approx(-200.0)
        # packaging on the 2 packs shipped
        assert s["Packaging (packs shipped)"] == pytest.approx(-20.0)
        revenue = (1200 + 216 - 216) + (-1199.0 + 182.90)
        assert s["Revenue (excl. GST)"] == pytest.approx(revenue)
        assert s["Contribution"] == pytest.approx(
            s["Contribution before GST credit"] + s["GST credit on fees (ITC)"])
        assert res["cash"]["deferred"] == pytest.approx(500.0)
        assert res["warnings"] == []

    def test_unmatched_and_uncosted_are_reported(self, service, engine):
        with engine.begin() as conn:
            conn.execute(text("UPDATE product_master SET total_unit_cogs = 0 WHERE sku = 'CONT-777'"))
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        service.import_file(make_flipkart_xlsx(), "f.xlsx", "account2")
        res = service.monthly_pnl("2026-09")
        text_ = " ".join(res["warnings"])
        assert "not in the catalog" in text_          # CAPS-PARA-WHITE has no Flipkart listing
        assert "no landed cost" in text_              # CONT-777 cost is 0
        assert list(res["unmatched"]["sku_raw"]) == ["CAPS-PARA-WHITE"]
        # suggested from the same SKU's Amazon listing
        assert res["unmatched"]["suggested_pack"].iloc[0] == "CAP-PARA-WHITE-PK1"

    def test_matching_a_sku_creates_listing_and_costs_it(self, service):
        service.import_file(make_flipkart_xlsx(), "f.xlsx", "account2")
        service.map_unmatched_sku("CAPS-PARA-WHITE", "flipkart", "account2", "CAP-PARA-WHITE-PK1")
        res = service.monthly_pnl("2026-09", platform="flipkart")
        assert res["unmatched"].empty
        # 2 sold, 1 returned -> 1 kept x Rs50; Flipkart GST from catalog (5%)
        assert res["summary"].loc["Goods cost (units kept)", "Total"] == pytest.approx(-50.0)
        assert not any("assumed" in w for w in res["warnings"])

    def test_per_sku_includes_shared_charges(self, service):
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        res = service.monthly_pnl("2026-09")
        assert res["by_sku"]["contribution"].sum() == pytest.approx(
            res["summary"].loc["Contribution", "Total"], abs=0.05)

    def test_filters_by_account(self, service):
        service.import_file(make_amazon_csv(), "a.csv", "account2")
        assert service.monthly_pnl("2026-09", account_id="account1") == {"empty": True}

    def test_word_match_suggestion(self, service, engine):
        with engine.connect() as conn:
            assert service.suggest_pack(conn, "L-CONT-777-PO2", "meesho") == "CONT-777-PO2"
